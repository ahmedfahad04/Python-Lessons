from openpyxl import load_workbook
from openpyxl.chart import BarChart, Reference

# load and activate the workbook
workbook = load_workbook(filename="PA3.6.xlsx")
sheet = workbook.active

# create Bar Chart
chart = BarChart()
chart.type = "col"
chart.style = 10
chart.title = "Player Results"
chart.x_axis.title = "Names"
chart.y_axis.title = "Scores"

# create data and title reference
data = Reference(sheet, min_col=2, min_row=1, max_row=7, max_col=2)
title = Reference(sheet, min_col=1, min_row=1, max_row=7)

# add data and title reference to chart
chart.add_data(data, titles_from_data=True)
chart.set_categories(title)
chart.shape = 4

# insert the chart into the sheet
sheet.add_chart(chart, "A10")
workbook.save("PA3.6_chart.xlsx")